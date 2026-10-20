"""Exportação CSV/Excel do razão e dos demonstrativos."""
from __future__ import annotations

import csv
import io

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select

from app.accounting.balance import current_position
from app.accounting.statements import dfc, dre, month_list
from app.database.session import session_scope
from app.models import Account, Category, EventType, FinancialEvent, Item, Institution, Merchant, Transaction

HEADER = PatternFill("solid", fgColor="111827")
BRL = '"R$" #,##0.00;[Red]-"R$" #,##0.00'


def ledger_rows(s) -> tuple[list[str], list[list]]:
    cats = {c.id: c for c in s.scalars(select(Category))}
    accs = {a.id: a for a in s.scalars(select(Account))}
    insts = {i.id: i.name for i in s.scalars(select(Institution))}
    items = {i.id: insts.get(i.institution_id, "") for i in s.scalars(select(Item))}
    merchants = {m.id: m.canonical_name for m in s.scalars(select(Merchant))}
    head = ["Data competência", "Data caixa", "Descrição", "Contraparte", "Comerciante", "Conta", "Instituição",
            "Tipo de evento", "Categoria", "Categoria-mãe", "Valor", "Regra", "Origem da regra", "Confiança",
            "Inferido", "Pendente", "Revisar", "Nota", "ID transação", "Categoria Pluggy"]
    tx = {t.id: t for t in s.scalars(select(Transaction))}
    rows = []
    for e in s.scalars(select(FinancialEvent).where(FinancialEvent.is_removed.is_(False))
                       .order_by(FinancialEvent.competence_date.desc())):
        c = cats.get(e.category_id)
        parent = cats.get(c.parent_id).name if c and c.parent_id in cats else (c.name if c else "")
        a = accs.get(e.account_id)
        t = tx.get(e.source_id)
        rows.append([e.competence_date, e.cash_date, e.description, e.counterparty, merchants.get(e.merchant_id),
                     a.name if a else "", items.get(a.item_id) if a else "", EventType.LABELS.get(e.event_type, e.event_type),
                     c.name if c else e.category_id, parent, e.amount, e.rule_id, e.rule_source, e.confidence,
                     "sim" if e.is_inferred else "não", "sim" if e.is_pending else "não",
                     "sim" if e.needs_review else "não", e.note, e.source_id, t.pluggy_category if t else None])
    return head, rows


def export_csv() -> bytes:
    with session_scope() as s:
        head, rows = ledger_rows(s)
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";")
    w.writerow(head)
    for r in rows:
        w.writerow([f"{v:.2f}".replace(".", ",") if isinstance(v, float) else v for v in r])
    return ("﻿" + buf.getvalue()).encode("utf-8")


def _sheet(wb: Workbook, title: str, head: list[str], rows: list[list], money: set[int] = frozenset()) -> None:
    ws = wb.create_sheet(title[:31])
    ws.append(head)
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), HEADER
    for r in rows:
        ws.append(r)
    for idx in money:
        for col in ws.iter_cols(min_col=idx, max_col=idx, min_row=2):
            for cell in col:
                cell.number_format = BRL
    for i, h in enumerate(head, 1):
        ws.column_dimensions[get_column_letter(i)].width = max(12, min(45, len(str(h)) + 6))
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions


def export_xlsx() -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    with session_scope() as s:
        head, rows = ledger_rows(s)
        _sheet(wb, "Razão", head, rows, {11})
        months = month_list(n=12)
        d = dre(s, months)
        drows = []
        for r in d["rows"]:
            drows.append([("= " if r["subtotal"] else "") + r["label"], *[r["values"][m] for m in months], r["total"]])
            for ch in r["children"]:
                drows.append(["    " + ch["label"], *[ch["values"][m] for m in months], ch["total"]])
        drows.append(["Taxa de poupança", *[d["savings_rate"][m] for m in months], None])
        if d.get("valuation"):
            drows.append(["Valorização da carteira (estim.)", *[d["valuation"]["values"].get(m) for m in months], None])
        _sheet(wb, "DRE", ["Linha", *months, "Total"], drows, set(range(2, len(months) + 3)))
        f = dfc(s, months)
        frows = []
        for sec in f["sections"]:
            frows.append([sec["label"], *[sec["values"][m] for m in months], sec["total"]])
            for ln in sec["lines"]:
                frows.append(["    " + ln["label"], *[ln["values"][m] for m in months], ln["total"]])
        frows += [["Variação de caixa", *[f["net_change"][m] for m in months], None],
                  ["Saldo inicial", *[f["opening_cash"][m] for m in months], None],
                  ["Saldo final", *[f["closing_cash"][m] for m in months], None]]
        _sheet(wb, "DFC", ["Linha", *months, "Total"], frows, set(range(2, len(months) + 3)))
        p = current_position(s)
        t = p["totals"]
        _sheet(wb, "Balanço", ["Item", "Valor"], [
            ["Caixa", t["cash"]], ["Investimentos", t["investments"]], ["Ativos manuais", t["manual_assets"]],
            ["= Ativo", t["assets"]], ["Cartão – fatura atual", t["card_current"]],
            ["Cartão – parcelas/compras futuras", t["card_future"]], ["Empréstimos", t["loans"]],
            ["Passivos manuais", t["manual_liabilities"]], ["= Passivo", t["liabilities"]],
            ["= Patrimônio líquido", t["net_worth"]], *[["Premissa", a] for a in p["assumptions"]]], {2})
        _sheet(wb, "Carteira", ["Ativo", "Código", "Classe", "Indexador", "Instituição", "Emissor", "Saldo", "Custo",
                                "Resultado", "Retorno", "% CDI", "Vencimento", "Peso"],
               [[x["name"], x["code"], x["asset_class"], x["indexer"], x["institution"], x["issuer"], x["balance"],
                 x["cost"], x["profit"], x["return_total"], x["pct_cdi"], x["due_date"], x["weight"]]
                for x in p["positions"]], {7, 8, 9})
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
